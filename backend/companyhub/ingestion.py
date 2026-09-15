"""Read-only workbook extraction; import immutable, auditable versions into isolated TEST."""
import hashlib
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from django.db import transaction

from .models import EmployeeMaster, PayrollRecord, SourceSnapshot


def scalar(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def extract_workbook(path):
    # Import dependency only for the operator command, not for serving the app.
    from openpyxl import load_workbook
    formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    cached = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    sheets, issues = [], []
    try:
        if len(formulas.sheetnames) > 60:
            raise ValueError("Nguồn có quá nhiều trang tính; cần xử lý thành đợt riêng.")
        for ws in formulas:
            if ws.max_row is None or ws.max_column is None:
                ws.calculate_dimension(force=True)
            if ws.max_row > 10000 or ws.max_column > 200:
                raise ValueError("Trang tính vượt giới hạn nhập an toàn của đợt này.")
            rows = []
            for row_number, (original, saved) in enumerate(zip(ws.iter_rows(), cached[ws.title].iter_rows()), start=1):
                cells = {}
                for cell, result in zip(original, saved):
                    if cell.value is None:
                        continue
                    item = {"value": scalar(result.value), "format": cell.number_format}
                    if cell.data_type == "f":
                        item["formula"] = cell.value
                        if result.value is None:
                            item["missing_cached_value"] = True
                        if "#REF!" in cell.value:
                            item["error"] = "#REF!"
                    if result.data_type == "e":
                        item["error"] = result.value
                    if item.get("error"):
                        issues.append({"sheet": ws.title, "cell": cell.coordinate, "message": f"Lỗi nguồn: {item['error']}"})
                    cells[cell.column_letter] = item
                if cells:
                    rows.append({"row": row_number, "cells": cells})
            sheets.append({"title": ws.title, "rows": rows, "row_count": ws.max_row, "column_count": ws.max_column})
    finally:
        formulas.close()
        cached.close()
    return {"kind": "workbook", "sheets": sheets}, issues


def cell_value(row, column):
    return row.get("cells", {}).get(column, {}).get("value")


def label(row, column):
    value = cell_value(row, column)
    return "" if value is None else str(value).strip()


def roster_rows(content):
    sheet = next(s for s in content["sheets"] if s["title"] == "THÔNG TIN NV")
    header = next(r for r in sheet["rows"] if r["row"] == 5)
    if label(header, "C") != "MSNV" or label(header, "D") != "HỌ VÀ TÊN":
        raise ValueError("Cấu trúc danh sách nhân viên thay đổi; dừng để tránh nhập nhầm cột.")
    records = []
    for row in sheet["rows"]:
        if row["row"] < 6 or not label(row, "D"):
            continue
        if not label(row, "C"):
            raise ValueError("Có dòng có tên nhưng thiếu mã: cần xem lại trước khi nhập.")
        issues = []
        company = label(row, "A")
        if not company:
            issues.append("Chưa có công ty trong nguồn; không tự điền từ dòng phía trên.")
        elif company not in {"LIVI", "AVI"}:
            issues.append("Chưa xác minh ánh xạ mã công ty nguồn với công ty trong app.")
        marker = label(row, "E").lower()
        resigned = cell_value(row, "CA") is not None
        status = "source_active" if marker == "x" else "source_resigned" if resigned else "unconfirmed"
        if marker == "x" and resigned:
            status = "conflict"
            issues.append("Vừa đánh dấu còn làm vừa có ngày nghỉ việc; cần xác minh.")
        if status == "unconfirmed":
            issues.append("Ô Còn làm trống; chưa kết luận đã nghỉ việc.")
        fields = [{"column": col, "label": label(header, col) or f"Cột {col}", **row["cells"].get(col, {"value": None})}
                  for col in header["cells"]]
        records.append({"source_row": row["row"], "code": label(row, "C"), "name": label(row, "D"),
                        "source_company": company, "department": label(row, "H"), "job_title": label(row, "I"),
                        "status": status, "fields": fields, "issues": issues})
    counts = Counter((r["source_company"], r["code"]) for r in records)
    for record in records:
        if counts[(record["source_company"], record["code"])] > 1:
            record["issues"].append("Mã nhân viên trùng trong cùng công ty nguồn; giữ từng dòng, không tự gộp.")
    return records


def payroll_rows(content, period):
    title = {"2026-05": "LƯƠNG T05.2026", "2026-06": "LƯƠNG T06.2026"}.get(period)
    if not title:
        raise ValueError("Kỳ lương chưa được kiểm chứng cấu trúc nhập.")
    sheet = next(s for s in content["sheets"] if s["title"] == title)
    header_number = 5 if period == "2026-05" else 1
    header = next(r for r in sheet["rows"] if r["row"] == header_number)
    if label(header, "B") != "MÃ SỐ NV" or label(header, "BC") != "THỰC LĨNH":
        raise ValueError("Cấu trúc lương thay đổi; không nhập theo vị trí cũ.")
    records = []
    for row in sheet["rows"]:
        if row["row"] <= header_number or not label(row, "B") or not label(row, "C"):
            continue
        issues = []
        cells = row["cells"]
        if "#REF!" in str(cells.get("AY", {}).get("formula", "")):
            issues.append("Ô thu nhập tính thuế mất tham chiếu #REF!.")
        if "AW" in str(cells.get("AZ", {}).get("formula", "")):
            issues.append("Công thức thuế tham chiếu AW (số người phụ thuộc), cần xác minh.")
        net_formula = str(cells.get("BC", {}).get("formula", ""))
        if net_formula and "AZ" not in net_formula:
            issues.append("Công thức thực lĩnh không trừ ô thuế AZ; chưa đủ căn cứ chốt.")
        gross = str(cells.get("AT", {}).get("formula", ""))
        if gross and "AM" not in gross:
            issues.append("Tổng thu nhập không cộng ô tổng tiền tăng ca AM; cần đối chiếu cách xử lý tăng ca.")
        records.append({"source_row": row["row"], "code": label(row, "B"), "name": label(row, "C"),
                        "department": label(row, "D"), "cells": cells, "issues": issues})
    return records


def prepare(path, metadata):
    path = Path(path)
    if path.stat().st_size > 30 * 1024 * 1024:
        raise ValueError("Tệp vượt giới hạn 30 MB của đợt nhập.")
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if path.suffix.lower() == ".xlsx":
        content, issues = extract_workbook(path)
    elif path.suffix.lower() == ".md":
        content, issues = {"kind": "document", "text": path.read_text(encoding="utf-8")}, []
    else:
        raise ValueError("Chỉ nhận bản xuất XLSX hoặc văn bản Markdown đã kiểm tra.")
    employees = roster_rows(content) if metadata["category"] == "employees" else []
    payroll = payroll_rows(content, metadata["period"]) if metadata["category"] == "payroll" else []
    return {"source": {**metadata, "digest": digest, "content": content, "issues": issues}, "employees": employees, "payroll": payroll}


@transaction.atomic
def commit_import(prepared):
    data = prepared["source"]
    source, created = SourceSnapshot.objects.get_or_create(source_id=data["source_id"], digest=data["digest"], defaults=data)
    if created:
        EmployeeMaster.objects.bulk_create([EmployeeMaster(source=source, **record) for record in prepared["employees"]])
        PayrollRecord.objects.bulk_create([PayrollRecord(source=source, **record) for record in prepared["payroll"]])
    return source, created
